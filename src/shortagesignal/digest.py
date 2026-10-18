"""
Day 12: the weekly digest email and the Excel list of items that need a look.

  python -m shortagesignal.digest            build the digest into digest/ to look at
  python -m shortagesignal.digest --send     also email it (needs the DIGEST_ settings)

The weekly workflow (.github/workflows/weekly.yml) runs this every Monday, after the daily
run, on the data branch's files. The email only goes out when these settings exist:
  DIGEST_SMTP_USER, DIGEST_SMTP_PASSWORD   the sending account (for Gmail, an app password)
  DIGEST_TO                                where to send it (defaults to DIGEST_SMTP_USER)
  DIGEST_SMTP_HOST, DIGEST_SMTP_PORT       default smtp.gmail.com and 587
Without them, the digest still shows on the run page and both files are saved as artifacts.

Every number in the digest is counted from the pipeline's tables. No AI writes it, so each
line can be traced back to the data.
"""
import argparse
import csv
import html
import os
import smtplib
import ssl
import sys
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill

from shortagesignal.config import (
    EXTRACT_REVIEW_PATH,
    HISTORY_PATH,
    LINKS_PATH,
    RECALLS_PATH,
    RISK_PATH,
    ROOT,
    RUN_LOG_PATH,
    TAG_REVIEW_PATH,
    WEEKLY_PATH,
)

DASHBOARD_URL = "https://shortagesignal.streamlit.app"
REPO_URL = "https://github.com/Showrish/shortagesignal"
FLAG_COLUMNS = ["flag", "points", "reported", "recall_number", "firm", "drug", "root_cause",
                "reasons", "class", "product"]
SHORT_COLUMNS = ["recall_number", "firm", "drug", "shortage_listing", "match_level"]


# ---------------------------------------------------------------- the facts

def _rows(con, sql, params=None):
    result = con.execute(sql, params or [])
    columns = [d[0] for d in result.description]
    return [dict(zip(columns, row)) for row in result.fetchall()]


def _read_csv(path):
    if not path.exists():
        return []
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def gather(today, paths=None):
    """Everything the digest says, counted from the pipeline's files."""
    p = {"risk": RISK_PATH, "recalls": RECALLS_PATH, "weekly": WEEKLY_PATH,
         "history": HISTORY_PATH, "links": LINKS_PATH, "tag_review": TAG_REVIEW_PATH,
         "extract_review": EXTRACT_REVIEW_PATH, "run_log": RUN_LOG_PATH, **(paths or {})}
    for name in ("risk", "recalls", "weekly", "history", "links"):
        if not p[name].exists():
            sys.exit(f"{p[name]} is missing. Copy the data branch in first: "
                     "git fetch origin data, then git archive origin/data | tar -x -C data")
    f = {name: path.as_posix() for name, path in p.items()}
    con = duckdb.connect()
    since = today - timedelta(days=7)

    this_week = _rows(con, f"""
        SELECT k.flag, k.points,
               strftime(TRY_STRPTIME(k.report_date, '%Y%m%d'), '%Y-%m-%d') AS reported,
               k.recall_number, r.recalling_firm AS firm, k.drug_key AS drug, k.root_cause,
               k.reasons, k.recall_class AS class, r.product_description AS product,
               r.event_id
        FROM read_parquet('{f['risk']}') AS k
        INNER JOIN read_parquet('{f['recalls']}') AS r ON r.recall_number = k.recall_number
        WHERE TRY_STRPTIME(k.report_date, '%Y%m%d') >= ?
        ORDER BY k.points DESC, reported DESC, k.recall_number
    """, [since])
    numbers = [r["recall_number"] for r in this_week]

    in_shortage = []
    if numbers:
        in_shortage = _rows(con, f"""
            SELECT DISTINCT l.recall_number, r.recalling_firm AS firm,
                   l.shortage_generic_name AS shortage_listing, l.match_level
            FROM read_parquet('{f['links']}') AS l
            INNER JOIN read_parquet('{f['recalls']}') AS r ON r.recall_number = l.recall_number
            WHERE l.shortage_list_type = 'shortage' AND lower(l.shortage_status) = 'current'
              AND l.recall_number IN ({", ".join("?" for _ in numbers)})
            ORDER BY l.recall_number
        """, numbers)
        drugs = {r["recall_number"]: r["drug"] for r in this_week}
        for r in in_shortage:
            r["drug"] = drugs.get(r["recall_number"])

    week = _rows(con, f"SELECT * FROM read_parquet('{f['weekly']}') ORDER BY week DESC LIMIT 1")
    recent_signals = con.execute(f"""
        SELECT COUNT(*) FROM (SELECT signal FROM read_parquet('{f['weekly']}')
                              ORDER BY week DESC LIMIT 13) WHERE signal <> ''
    """).fetchone()[0]

    # Listings that appeared or left this week. The first day of snapshots is skipped,
    # since every listing was "new" that day.
    added, left = con.execute(f"""
        WITH h AS (SELECT * FROM read_parquet('{f['history']}')),
             bounds AS (SELECT min(first_seen) AS start, max(last_seen) AS latest FROM h)
        SELECT COUNT(*) FILTER (WHERE first_seen > start AND first_seen >= ?),
               COUNT(*) FILTER (WHERE last_seen < latest AND last_seen >= ?)
        FROM h, bounds
    """, [since, since]).fetchone()

    runs = [r for r in _read_csv(p["run_log"]) if r.get("run_at_utc", "")[:10] >= since.isoformat()]
    tag_review = _read_csv(p["tag_review"])
    # Only this week's recalls. The full backlog stays in extract_review.csv on the data branch.
    extract_review = [r for r in _read_csv(p["extract_review"])
                      if r.get("problem") in ("not a known drug name", "low confidence")
                      and r.get("recall_number") in set(numbers)]
    return {
        "today": today, "since": since,
        "recalls": len(this_week), "events": len({r["event_id"] for r in this_week}),
        "flagged": [r for r in this_week if r["flag"] in ("high", "watch")],
        "high": sum(1 for r in this_week if r["flag"] == "high"),
        "watch": sum(1 for r in this_week if r["flag"] == "watch"),
        "in_shortage": in_shortage,
        "week": week[0] if week else None, "recent_signals": recent_signals,
        "listings_added": added, "listings_left": left,
        "runs": len(runs), "runs_ok": sum(1 for r in runs if r.get("status") == "success"),
        "warnings": int(sum(_number(r.get("checks_warned")) for r in runs)),
        "tag_review": tag_review,
        "tag_review_recalls": int(sum(_number(r.get("recalls")) for r in tag_review)),
        "extract_review": extract_review,
    }


# ---------------------------------------------------------------- the words

def _s(n, word, plural=None):
    return f"{n:,} {word if n == 1 else plural or word + 's'}"


def compose(facts):
    """(subject, sections). Each section is (heading, paragraphs, table or None)."""
    t = facts["today"]
    subject = (f"ShortageSignal weekly, {t:%b} {t.day}: {facts['high']} high and "
               f"{facts['watch']} watch recalls" if facts["recalls"]
               else f"ShortageSignal weekly, {t:%b} {t.day}: no new recalls")
    sections = []

    w = facts["week"]
    if facts["recalls"]:
        line = (f"In the 7 days to {t:%B} {t.day}, FDA reported {_s(facts['recalls'], 'recall')} "
                f"from {_s(facts['events'], 'recall event')}.")
    else:
        line = f"FDA reported no new recalls in the 7 days to {t:%B} {t.day}."
    if w and w.get("ucl") is not None:
        place = {"above": "above", "below": "below"}.get(w.get("signal"), "inside")
        line += (f" The last full week, starting {w['week']}, had {_s(int(w['events']), 'recall event')}, "
                 f"{place} the normal range of {w['lcl']:.0f} to {w['ucl']:.0f}.")
        if facts["recent_signals"]:
            n = facts["recent_signals"]
            line += f" {_s(n, 'week')} in the last 13 {'was' if n == 1 else 'were'} outside it."
    sections.append(("This week", [line], None))

    if facts["flagged"]:
        para = (f"{_s(facts['high'], 'recall')} flagged high and {facts['watch']:,} watch. "
                "Points come from a sterility problem (2), a drug only 1 or 2 companies make "
                "(2, or 1 for 3 to 5 companies) and an injectable (1). 4 or more is high, 2 or 3 "
                "is watch. On past recalls, high flags were followed by a shortage about 4 times "
                "as often as average.")
        table = (["flag", "drug", "firm", "root cause", "reasons"],
                 [[r["flag"], r["drug"] or "(not named)", r["firm"], r["root_cause"], r["reasons"]]
                  for r in facts["flagged"][:10]])
        sections.append(("Recalls worth watching", [para], table))
    else:
        sections.append(("Recalls worth watching", ["No recall this week was flagged high or watch."], None))

    if facts["in_shortage"]:
        para = (f"{_s(len({r['recall_number'] for r in facts['in_shortage']}), 'recall')} this week "
                "covered a drug FDA lists as in shortage right now. A recall like that can make "
                "an existing shortage worse.")
        table = (["recall", "drug", "firm", "shortage listing", "match"],
                 [[r["recall_number"], r["drug"] or "", r["firm"], r["shortage_listing"],
                   r["match_level"].replace("_", " ")] for r in facts["in_shortage"][:10]])
        sections.append(("Already in shortage", [para], table))

    sections.append(("FDA's shortage list", [
        f"{_s(facts['listings_added'], 'listing')} added and {facts['listings_left']:,} removed "
        "this week, across the shortage and discontinuation lists."], None))

    extract = len(facts["extract_review"])
    sections.append(("Needs a look", [
        f"{_s(len(facts['tag_review']), 'unsure root-cause tag')} covering "
        f"{_s(facts['tag_review_recalls'], 'recall')} in total, and {_s(extract, 'drug name')} "
        "Claude suggested for this week's recalls that weren't used. Both lists are in the "
        "attached Excel file."], None))

    # Only what a reader needs to trust this week's numbers. Costs and other developer
    # details stay on the run page and in run_log.csv.
    health = (f"{facts['runs_ok']} of {_s(facts['runs'], 'pipeline run')} worked this week"
              + (f", with {_s(facts['warnings'], 'data check warning')}." if facts["warnings"]
                 else ", and every data check passed."))
    sections.append(("Data behind this digest", [health], None))
    return subject, sections


def to_text(subject, sections):
    lines = [subject, ""]
    for heading, paragraphs, table in sections:
        lines += [heading.upper(), *paragraphs]
        if table:
            for row in table[1]:
                lines.append("  - " + " | ".join(str(v) for v in row))
        lines.append("")
    lines += [f"Dashboard: {DASHBOARD_URL}", f"Code: {REPO_URL}"]
    return "\n".join(lines)


def to_markdown(subject, sections):
    lines = [f"# {subject}", ""]
    for heading, paragraphs, table in sections:
        lines += [f"## {heading}", "", *[p + "\n" for p in paragraphs]]
        if table:
            headers, rows = table
            lines += ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
            lines += ["| " + " | ".join(str(v).replace("|", "/") for v in row) + " |" for row in rows]
            lines.append("")
    lines.append(f"[Dashboard]({DASHBOARD_URL}) · [Code]({REPO_URL})")
    return "\n".join(lines)


def to_html(subject, sections):
    e = html.escape
    parts = [f"<h2>{e(subject)}</h2>"]
    for heading, paragraphs, table in sections:
        parts.append(f"<h3>{e(heading)}</h3>")
        parts += [f"<p>{e(p)}</p>" for p in paragraphs]
        if table:
            headers, rows = table
            parts.append('<table cellpadding="6" style="border-collapse:collapse">')
            parts.append("<tr>" + "".join(f'<th align="left" style="border-bottom:1px solid #999">'
                                          f"{e(h)}</th>" for h in headers) + "</tr>")
            parts += ["<tr>" + "".join(f"<td>{e(str(v))}</td>" for v in row) + "</tr>" for row in rows]
            parts.append("</table>")
    parts.append(f'<p><a href="{DASHBOARD_URL}">Dashboard</a> · <a href="{REPO_URL}">Code</a></p>')
    return "<html><body style=\"font-family:sans-serif\">" + "\n".join(parts) + "</body></html>"


# ---------------------------------------------------------------- the Excel list

def _cell(value):
    """Some FDA text has hidden control characters, which Excel can't store. Drop them."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def write_workbook(facts, path):
    """The items that need a look, one sheet each, ready to filter."""
    wb = Workbook()
    wb.remove(wb.active)
    sheets = [
        ("Flagged this week", FLAG_COLUMNS, facts["flagged"]),
        ("Already in shortage", SHORT_COLUMNS, facts["in_shortage"]),
        ("Unsure tags", ["reason_for_recall", "recalls", "root_cause", "confidence", "why",
                         "example_recall"], facts["tag_review"]),
        ("Drug names not used", ["recall_number", "recalling_firm", "product_description",
                                 "ingredients", "confidence", "problem"], facts["extract_review"]),
    ]
    header_fill = PatternFill("solid", fgColor="DDDDDD")
    for title, columns, rows in sheets:
        ws = wb.create_sheet(title)
        ws.append(columns)
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True), header_fill
        for row in rows:
            ws.append([_cell(row.get(c)) for c in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, column in enumerate(columns, start=1):
            longest = max([len(str(column))] + [len(str(r.get(column) or "")) for r in rows[:200]])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(12, longest + 2), 60)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------- the email

def email_settings():
    """The DIGEST_ settings, or None when the email isn't set up."""
    user, password = os.getenv("DIGEST_SMTP_USER"), os.getenv("DIGEST_SMTP_PASSWORD")
    if not user or not password:
        return None
    return {"user": user, "password": password, "to": os.getenv("DIGEST_TO") or user,
            "host": os.getenv("DIGEST_SMTP_HOST") or "smtp.gmail.com",
            "port": int(os.getenv("DIGEST_SMTP_PORT") or 587)}


def build_message(subject, text, page, attachment, settings):
    message = EmailMessage()
    message["Subject"], message["From"], message["To"] = subject, settings["user"], settings["to"]
    message.set_content(text)
    message.add_alternative(page, subtype="html")
    message.add_attachment(
        attachment.read_bytes(), maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=attachment.name)
    return message


def send(message, settings):
    with smtplib.SMTP(settings["host"], settings["port"], timeout=60) as server:
        server.starttls(context=ssl.create_default_context())
        server.login(settings["user"], settings["password"])
        server.send_message(message)


# ---------------------------------------------------------------- run

def main(argv=None):
    parser = argparse.ArgumentParser(description="Build the weekly digest, and email it with --send.")
    parser.add_argument("--send", action="store_true", help="email it, if the DIGEST_ settings exist")
    parser.add_argument("--out", default=str(ROOT / "digest"), help="where to save the files")
    args = parser.parse_args(sys.argv[1:] if argv is None else argv)

    today = datetime.now(timezone.utc).date()
    facts = gather(today)
    subject, sections = compose(facts)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    workbook = write_workbook(facts, out / f"needs_a_look_{today.isoformat()}.xlsx")
    page, markdown = to_html(subject, sections), to_markdown(subject, sections)
    (out / "digest.html").write_text(page, encoding="utf-8")
    (out / "digest.md").write_text(markdown, encoding="utf-8")
    print(to_text(subject, sections))
    if os.getenv("GITHUB_STEP_SUMMARY"):
        with open(os.environ["GITHUB_STEP_SUMMARY"], "a", encoding="utf-8") as fh:
            fh.write(markdown + "\n")

    if args.send:
        settings = email_settings()
        if settings is None:
            print("\nEmail not set up (no DIGEST_SMTP_USER and DIGEST_SMTP_PASSWORD), so nothing was sent.")
        else:
            send(build_message(subject, to_text(subject, sections), page, workbook, settings), settings)
            print(f"\nEmailed to {settings['to']}")
    print(f"Saved {out}")


if __name__ == "__main__":
    main()
