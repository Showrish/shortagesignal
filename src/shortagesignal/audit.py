"""
Day 7: the hand audit.

  python -m shortagesignal.audit sample   picks random recalls (200, or --size N) into audit/audit_sheet.xlsx
  python -m shortagesignal.audit score    reads your filled-in sheet and scores the pipeline

NDC checking is mostly automatic. The sheet lists the codes the pipeline found in the
description, plus any other code-like numbers it didn't use (leftover_numbers). When there
are no leftovers, ndc_ok is filled in for you. You only look at rows with leftovers.

The sheet shows what the pipeline concluded for each recall in grey columns, next to
yellow columns you fill in. Scoring turns your answers into:
  audit/answer_key.csv     your verdicts and root-cause labels, the answer key for Day 8
  audit/audit_results.md   how often the pipeline was right, written as "caught X of Y"
"""
import argparse
import csv
import re
import sys
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from shortagesignal.checks import RECALL_FIELDS
from shortagesignal.config import LINKS_PATH, RECALL_DRUGS_PATH, RECALL_IDS_PATH, RECALLS_PATH, ROOT
from shortagesignal.ndc import DIGITS_RE, PACKAGE_RE, PRODUCT_RE

AUDIT_DIR = ROOT / "audit"
SHEET_PATH = AUDIT_DIR / "audit_sheet.xlsx"
SAMPLE_SIZE = 200
SEED = "shortagesignal-audit-1"   # same seed, same 200 recalls, every time

ROOT_CAUSES = {
    "CGMP": "The factory broke manufacturing rules (CGMP deviations), with no more specific cause given",
    "Sterility": "A sterile product can't be shown to be sterile, or a sterile product was contaminated",
    "Microbial contamination": "Bacteria, mold or yeast found in a product that isn't meant to be sterile",
    "Chemical or foreign contamination": "A wrong chemical, another drug (like penicillin cross-contamination), or foreign material",
    "Particulates": "Visible or tiny particles in an injectable",
    "Impurities": "Failed impurity or degradation limits, like the NDMA impurities in valsartan",
    "Potency": "Too weak (subpotent) or too strong (superpotent)",
    "Failed specification": "Failed dissolution, content uniformity, pH or another test not covered above",
    "Labeling or packaging": "Wrong or missing label information, label mix-ups, the wrong product in the package",
    "Container or device defect": "Leaking, broken or defective containers, syringes, pumps or delivery systems",
    "Unapproved drug": "Marketed without an approved application (NDA or ANDA)",
    "Other": "Anything that doesn't fit. Explain it in notes",
}
DROPDOWNS = {
    "drug_ok": ["yes", "no", "can't tell"],
    "ndc_ok": ["yes", "missed some", "wrong", "no NDC in text"],
    "data_problem": ["none", "missing field", "unreadable text", "other"],
    "root_cause": list(ROOT_CAUSES),
}
PIPELINE_COLUMNS = ["recall_number", "report_date", "class", "firm", "product_description",
                    "reason_for_recall", "ndcs_in_text", "leftover_numbers", "drug_found",
                    "how_found", "linked_shortage", "fields_missing"]
YOUR_COLUMNS = ["drug_ok", "correct_drug", "ndc_ok", "data_problem", "root_cause", "notes"]
WIDTHS = {"product_description": 60, "reason_for_recall": 50, "ndcs_in_text": 22,
          "leftover_numbers": 18, "drug_found": 24,
          "firm": 24, "correct_drug": 22, "root_cause": 26, "notes": 30}
HOW_TO = [
    ("drug_ok", "yes if drug_found names the right drug. no if it's wrong, or if it's empty but the "
                "description clearly names the drug. can't tell if the description doesn't say"),
    ("correct_drug", "Only when drug_ok is no: the right generic name, like heparin sodium"),
    ("ndc_ok", "Already filled in when leftover_numbers is blank. When leftover_numbers has "
               "something, check it: if one of those is an NDC the pipeline missed, pick missed some. "
               "If they're lot numbers, dates or other numbers, pick yes, or no NDC in text when "
               "ndcs_in_text is empty. Pick wrong if ndcs_in_text has a code that isn't in the "
               "description. ndcs_in_text uses the 11-digit format, so 0409-2720-01 shows as "
               "00409-2720-01, and that's correct"),
    ("data_problem", "none, or missing field (a field like firm or reason is blank), "
                     "unreadable text, or other"),
    ("root_cause", "The one category that best fits reason_for_recall. See the list below"),
    ("notes", "Anything worth remembering. Optional"),
]
EXAMPLE = {
    "recall_number": "D-0123-2026", "report_date": "20260107", "class": "Class II",
    "firm": "Hospira, Inc.",
    "product_description": "Heparin Sodium Injection, USP, 1,000 USP Units/mL, 10 mL vial, NDC 0409-2720-01",
    "reason_for_recall": "Lack of Assurance of Sterility: a leak was found in the vial seal",
    "ndcs_in_text": "00409-2720-01", "leftover_numbers": "", "drug_found": "heparin sodium",
    "how_found": "ndc_directory",
    "linked_shortage": "package", "fields_missing": "",
    "drug_ok": "yes", "correct_drug": "", "ndc_ok": "yes", "data_problem": "none",
    "root_cause": "Sterility", "notes": "",
}

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
GREY = PatternFill("solid", start_color="E7E6E6")
YELLOW = PatternFill("solid", start_color="FFFF00")
WRAP = Alignment(wrap_text=True, vertical="top")


# ---------------------------------------------------------------- the sheet

def make_sheet(rows, path):
    """Write the audit workbook: an Audit tab to fill in, and a How to fill tab."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit"
    columns = PIPELINE_COLUMNS + YOUR_COLUMNS
    ws.append(columns)
    for row in rows:
        ws.append([row.get(c, "") for c in PIPELINE_COLUMNS + YOUR_COLUMNS])

    for i, name in enumerate(columns, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = WIDTHS.get(name, 15)
        mine = name in YOUR_COLUMNS
        for cell in ws[letter]:
            cell.font = FONT
            cell.alignment = WRAP
            if mine:
                cell.fill = YELLOW
        ws.cell(row=1, column=i).font = BOLD
        if not mine:
            ws.cell(row=1, column=i).fill = GREY
        if name in DROPDOWNS:
            rule = DataValidation(type="list", formula1='"' + ",".join(DROPDOWNS[name]) + '"',
                                  allow_blank=True)
            rule.add(f"{letter}2:{letter}{len(rows) + 1}")
            ws.add_data_validation(rule)
    ws.freeze_panes = "B2"

    guide = wb.create_sheet("How to fill")
    guide.column_dimensions["A"].width = 32
    guide.column_dimensions["B"].width = 100
    lines = [("Grey columns", "What the pipeline concluded. Don't edit these"),
             ("Yellow columns", "Yours to fill in, one row per recall. Most use dropdowns"),
             ("", "")]
    lines += HOW_TO + [("", ""), ("Root causes", "")] + list(ROOT_CAUSES.items())
    lines += [("", ""), ("Example row", "A filled-in row looks like this")]
    for label, text in lines:
        guide.append([label, text])
    guide.append(columns)
    guide.append([EXAMPLE[c] for c in columns])
    for row in guide.iter_rows():
        for cell in row:
            cell.font = FONT
            cell.alignment = WRAP
    for label in ("Grey columns", "Yellow columns", "Root causes", "Example row"):
        for row in guide.iter_rows(min_col=1, max_col=1):
            if row[0].value == label:
                row[0].font = BOLD
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# Anything that looks like a code: three groups (dashes, spaces or any dash character
# between them), two groups with a dash, or a run of 10 to 11 digits.
_SEP = r"[\s\-\u2010-\u2015]"
CODE_LIKE = re.compile(rf"(?<![\d-])\d{{4,5}}{_SEP}\d{{3,4}}{_SEP}\d{{1,2}}(?![\d-])"
                       r"|(?<!\d)\d{10,11}(?!\d)"
                       r"|(?<![\d-])\d{4,5}-\d{3,4}(?![\d-])")


def leftover_numbers(text):
    """Code-like numbers in a description that the NDC extractor didn't pick up."""
    if not text:
        return ""
    chars = list(text)
    for rx in (PACKAGE_RE, PRODUCT_RE, DIGITS_RE):
        for match in rx.finditer(text):
            start, end = match.span(1)
            chars[start:end] = " " * (end - start)
    found = []
    for match in CODE_LIKE.finditer("".join(chars)):
        value = match.group(0).strip()
        if value not in found:
            found.append(value)
    return ", ".join(found)


def prefill_ndc_ok(row):
    """ndc_ok is obvious when nothing code-like is left over, so fill it in."""
    if row.get("leftover_numbers"):
        return ""
    return "yes" if row.get("ndcs_in_text") else "no NDC in text"


def carry_over(rows, old_rows):
    """Keep any answers already filled in, matched by recall number."""
    old = {r.get("recall_number"): r for r in old_rows}
    for row in rows:
        previous = old.get(row["recall_number"], {})
        for column in YOUR_COLUMNS:
            if previous.get(column):
                row[column] = previous[column]
    return rows


def fields_missing(row):
    return ", ".join(f for f in RECALL_FIELDS if not str(row.get(f) or "").strip())


def sample_rows(con, n=SAMPLE_SIZE, seed=SEED):
    """n random recalls (the same ones every time), with what the pipeline concluded."""
    p = {name: path.as_posix() for name, path in {
        "recalls": RECALLS_PATH, "recall_ids": RECALL_IDS_PATH,
        "recall_drugs": RECALL_DRUGS_PATH, "links": LINKS_PATH}.items()}
    rel = con.sql(f"""
        WITH picked AS (
            SELECT * FROM read_parquet('{p['recalls']}')
            ORDER BY md5(recall_number || '{seed}')
            LIMIT {n}
        ),
        ndcs AS (
            SELECT recall_number,
                   string_agg(DISTINCT coalesce(ndc11, ndc9), ', ')
                       FILTER (WHERE source IN ('text', 'text_product', 'directory_lookup')) AS ndcs
            FROM read_parquet('{p['recall_ids']}') WHERE ndc9 IS NOT NULL
            GROUP BY recall_number
        ),
        best AS (
            SELECT recall_number, min(match_rank) AS best
            FROM read_parquet('{p['links']}') GROUP BY recall_number
        )
        SELECT p.*, coalesce(n.ndcs, '') AS ndcs_in_text,
               coalesce(d.drug_key, '') AS drug_found, coalesce(d.id_source, '') AS how_found,
               CASE b.best WHEN 1 THEN 'package' WHEN 2 THEN 'product'
                           WHEN 3 THEN 'drug and firm' WHEN 4 THEN 'drug' ELSE 'none' END
                   AS linked_shortage
        FROM picked p
        LEFT JOIN ndcs n ON n.recall_number = p.recall_number
        LEFT JOIN read_parquet('{p['recall_drugs']}') d ON d.recall_number = p.recall_number
        LEFT JOIN best b ON b.recall_number = p.recall_number
        ORDER BY p.report_date DESC
    """)
    rows = [dict(zip(rel.columns, values)) for values in rel.fetchall()]
    for row in rows:
        row["class"] = row.get("classification")
        row["firm"] = row.get("recalling_firm")
        row["fields_missing"] = fields_missing(row)
        row["leftover_numbers"] = leftover_numbers(row.get("product_description"))
        row["ndc_ok"] = prefill_ndc_ok(row)
    return rows


# ---------------------------------------------------------------- scoring

def _clean(value):
    return str(value).strip() if value is not None else ""


def read_sheet(path):
    ws = load_workbook(path, read_only=True, data_only=True)["Audit"]
    rows = ws.iter_rows(values_only=True)
    header = [_clean(h) for h in next(rows)]
    return [{h: _clean(v) for h, v in zip(header, values)}
            for values in rows if any(_clean(v) for v in values)]


def _rate(caught, total):
    return f"{caught} of {total} ({round(100 * caught / total)}%)" if total else "nothing to score yet"


def _as_date(text):
    """20260610 → 2026-06-10."""
    return f"{text[:4]}-{text[4:6]}-{text[6:8]}" if len(text) == 8 and text.isdigit() else text


def score(rows):
    """Turn the filled-in rows into results. Only rows with drug_ok filled count as audited."""
    audited = [r for r in rows if r.get("drug_ok")]
    drug_known = [r for r in audited if r["drug_ok"] in ("yes", "no")]
    named = [r for r in drug_known if r.get("drug_found")]
    right = [r for r in drug_known if r["drug_ok"] == "yes"]
    with_ndc = [r for r in audited if r.get("ndc_ok") in ("yes", "missed some")]
    all_ndcs = [r for r in with_ndc if r["ndc_ok"] == "yes"]
    false_codes = [r for r in audited if r.get("ndc_ok") == "wrong"]
    missing = [r for r in audited if r.get("data_problem") == "missing field"]
    flagged = [r for r in missing if r.get("fields_missing")]
    false_alarms = [r for r in audited if r.get("fields_missing") and r.get("data_problem") == "none"]
    by_hand = [r for r in audited if r.get("leftover_numbers") and r.get("ndc_ok")]
    dates = sorted(r["report_date"] for r in audited if r.get("report_date"))
    causes = {}
    for r in audited:
        if r.get("root_cause"):
            causes[r["root_cause"]] = causes.get(r["root_cause"], 0) + 1
    return {
        "audited": len(audited),
        "first_date": _as_date(dates[0]) if dates else "",
        "last_date": _as_date(dates[-1]) if dates else "",
        "drug_precision": _rate(len(right), len(named)),
        "drug_coverage": _rate(len(right), len(drug_known)),
        "ndc_capture": _rate(len(all_ndcs), len(with_ndc)),
        "false_codes": len(false_codes),
        "ndc_by_hand": len(by_hand),
        "missing_flagged": _rate(len(flagged), len(missing)),
        "false_alarms": len(false_alarms),
        "root_causes": dict(sorted(causes.items(), key=lambda kv: -kv[1])),
        "labeled": sum(causes.values()),
    }


def write_results(rows, result, out_dir, today):
    """Save the answer key and the results page."""
    out_dir.mkdir(parents=True, exist_ok=True)
    with open(out_dir / "answer_key.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["recall_number", "reason_for_recall"] + YOUR_COLUMNS,
                                extrasaction="ignore")
        writer.writeheader()
        writer.writerows(r for r in rows if r.get("drug_ok"))

    lines = [
        "# Hand audit results", "",
        f"Audited {result['audited']} of {len(rows)} sampled recalls, scored on {today}.",
        f"The audited recalls were reported between {result['first_date']} and "
        f"{result['last_date']}, so these results describe that period.", "",
        "| What the pipeline does | Caught |", "|---|---|",
        f"| Names the right drug, when it names one | {result['drug_precision']} |",
        f"| Finds the drug, for recalls where the drug is clear | {result['drug_coverage']} |",
        f"| Finds every NDC in the description | {result['ndc_capture']}, "
        f"plus {result['false_codes']} recalls with a wrong code. {result['ndc_by_hand']} rows "
        "had leftover code-like numbers and were checked by hand, the rest were clear |",
        f"| Flags recalls missing a required field | {result['missing_flagged']}, "
        f"plus {result['false_alarms']} false alarms |",
        "", f"## Root causes ({result['labeled']} labeled)", "",
        "| Root cause | Recalls |", "|---|---|",
    ]
    lines += [f"| {cause} | {n} |" for cause, n in result["root_causes"].items()]
    (out_dir / "audit_results.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return lines


# ---------------------------------------------------------------- command line

def main(argv=None):
    parser = argparse.ArgumentParser(prog="python -m shortagesignal.audit")
    parser.add_argument("step", choices=["sample", "score"])
    parser.add_argument("--force", action="store_true",
                        help="rebuild an existing sheet, keeping any answers already in it")
    parser.add_argument("--size", type=int, default=SAMPLE_SIZE, help="how many recalls to sample")
    args = parser.parse_args(argv)

    if args.step == "sample":
        if SHEET_PATH.exists() and not args.force:
            sys.exit(f"{SHEET_PATH} already exists, and it may have your answers in it. "
                     "Add --force to rebuild it. Answers already in it are kept.")
        import duckdb
        rows = sample_rows(duckdb.connect(), n=args.size)
        if SHEET_PATH.exists():
            rows = carry_over(rows, read_sheet(SHEET_PATH))
        make_sheet(rows, SHEET_PATH)
        print(f"Wrote {len(rows)} recalls to {SHEET_PATH}. Download it, fill in the yellow columns, "
              "then upload it back to the same place and run: python -m shortagesignal.audit score")
    else:
        rows = read_sheet(SHEET_PATH)
        result = score(rows)
        for line in write_results(rows, result, AUDIT_DIR, date.today().isoformat()):
            print(line)
        print(f"\nSaved {AUDIT_DIR / 'answer_key.csv'} and {AUDIT_DIR / 'audit_results.md'}")


if __name__ == "__main__":
    main()
