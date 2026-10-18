"""The weekly digest: what it counts, what it says, the Excel list and the email."""
import csv
from datetime import date

import duckdb
from openpyxl import load_workbook

from shortagesignal.digest import (
    build_message,
    compose,
    email_settings,
    gather,
    keeping_it_running,
    to_html,
    to_markdown,
    to_text,
    upkeep,
    write_workbook,
)

TODAY = date(2026, 10, 19)


def _parquet(con, path, columns, rows):
    values = ", ".join("(" + ", ".join(rows_value(v) for v in row) + ")" for row in rows)
    con.execute(f"COPY (SELECT * FROM (VALUES {values}) t({', '.join(columns)})) "
                f"TO '{path.as_posix()}' (FORMAT PARQUET)")


def rows_value(v):
    if v is None:
        return "NULL"
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, date):
        return f"DATE '{v.isoformat()}'"
    return "'" + str(v).replace("'", "''") + "'"


def _csv(path, rows):
    with open(path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _fixture(tmp_path):
    con = duckdb.connect()
    p = {name: tmp_path / f"{name}.parquet" for name in ("risk", "recalls", "weekly", "history", "links")}
    _parquet(con, p["risk"], ["recall_number", "report_date", "flag", "points", "drug_key",
                              "root_cause", "reasons", "recall_class"], [
        ("D-1", "20261014", "high", 5, "hypromellose", "Sterility", "sterility, 2 labelers", "Class II"),
        ("D-2", "20261014", "watch", 3, "dopamine hydrochloride", "Sterility", "sterility, injectable", "Class II"),
        ("D-3", "20261013", "low", 0, "aspirin", "Labeling or packaging", "", "Class III"),
        ("D-4", "20260901", "high", 5, "heparin sodium", "Sterility", "sterility, 1 labeler", "Class I"),
    ])
    _parquet(con, p["recalls"], ["recall_number", "recalling_firm", "product_description", "event_id"], [
        ("D-1", "Aurolab", "Eye drops", "E1"), ("D-2", "Aurolab", "Dopamine vial", "E1"),
        ("D-3", "Acme", "Aspirin tablets", "E2"), ("D-4", "Hospira", "Heparin injection", "E0"),
    ])
    _parquet(con, p["weekly"], ["week", "events", "recalls", "center", "ucl", "lcl", "signal"], [
        (date(2026, 9, 28), 20, 25, 6.0, 15.0, 0.0, "above"),
        (date(2026, 10, 5), 7, 9, 6.0, 15.0, 0.0, ""),
    ])
    _parquet(con, p["history"], ["shortage_key", "first_seen", "last_seen"], [
        ("k1", date(2026, 10, 11), date(2026, 10, 19)),   # there since the first snapshot
        ("k2", date(2026, 10, 14), date(2026, 10, 19)),   # added this week
        ("k3", date(2026, 10, 11), date(2026, 10, 15)),   # left this week
    ])
    _parquet(con, p["links"], ["recall_number", "shortage_list_type", "shortage_status",
                               "shortage_generic_name", "match_level"], [
        ("D-2", "shortage", "Current", "Dopamine HCl", "drug"),
        ("D-3", "shortage", "Resolved", "Aspirin", "drug"),          # not current
        ("D-4", "shortage", "Current", "Heparin", "package"),       # an old recall
    ])
    p["run_log"] = tmp_path / "run_log.csv"
    _csv(p["run_log"], [
        {"run_at_utc": "2026-10-01T12:00:00", "status": "success", "checks_warned": "0",
         "claude_cost_usd": "1.00", "extract_cost_usd": "0", "tag_backlog": "0",
         "extract_waiting": "0"},
        *[{"run_at_utc": f"2026-10-{day}T12:00:00", "status": "success", "checks_warned": "0",
           "claude_cost_usd": "0.01", "extract_cost_usd": "0.01", "tag_backlog": "2",
           "extract_waiting": "1"} for day in (13, 14, 15, 16, 17)],
        {"run_at_utc": "2026-10-18T12:00:00", "status": "failed", "checks_warned": "1",
         "claude_cost_usd": "", "extract_cost_usd": "", "tag_backlog": "", "extract_waiting": ""},
    ])
    p["tag_review"] = tmp_path / "tag_review.csv"
    _csv(p["tag_review"], [
        {"reason_for_recall": "Odd", "recalls": "3", "root_cause": "Other", "confidence": "medium",
         "why": "", "example_recall": "D-9"},
        {"reason_for_recall": "Odd too", "recalls": "1", "root_cause": "CGMP", "confidence": "low",
         "why": "", "example_recall": "D-8"},
    ])
    p["extract_review"] = tmp_path / "extract_review.csv"
    _csv(p["extract_review"], [
        {"recall_number": "D-3", "recalling_firm": "X", "product_description": "Tri-mix",
         "ingredients": "a, b", "confidence": "high", "problem": "not a known drug name"},
        {"recall_number": "D-6", "recalling_firm": "Y", "product_description": "Powder",
         "ingredients": "c", "confidence": "low", "problem": "low confidence"},
        {"recall_number": "D-5", "recalling_firm": "Z", "product_description": "Device",
         "ingredients": "", "confidence": "high", "problem": "no drug named"},
    ])
    return p


NO_SETTINGS = {"bought": None, "since": None, "key_expires": None}


def test_gather_counts_the_week(tmp_path):
    f = gather(TODAY, _fixture(tmp_path), NO_SETTINGS)
    assert f["waiting"] == 3                   # the latest run that worked: 2 tags and 1 drug name
    assert (f["recalls"], f["events"], f["high"], f["watch"]) == (3, 2, 1, 1)
    assert [r["recall_number"] for r in f["flagged"]] == ["D-1", "D-2"]        # highest points first
    assert [(r["recall_number"], r["drug"]) for r in f["in_shortage"]] == [("D-2", "dopamine hydrochloride")]
    assert (f["week"]["events"], f["recent_signals"]) == (7, 1)
    assert (f["listings_added"], f["listings_left"]) == (1, 1)
    assert (f["runs"], f["runs_ok"], f["warnings"]) == (6, 5, 1)
    # D-3 is this week's, D-6 isn't, and "no drug named" never counts
    assert (len(f["tag_review"]), f["tag_review_recalls"], len(f["extract_review"])) == (2, 4, 1)


def test_compose_says_it_in_plain_english(tmp_path):
    subject, sections = compose(gather(TODAY, _fixture(tmp_path), NO_SETTINGS))
    assert subject == "ShortageSignal weekly, Oct 19: 1 high and 1 watch recalls"
    text = to_text(subject, sections)
    assert "FDA reported 3 recalls from 2 recall events" in text
    assert "inside the normal range of 0 to 15" in text
    assert "1 recall this week covered a drug FDA lists as in shortage right now" in text
    assert "5 of 6 pipeline runs worked this week, with 1 data check warning." in text
    assert "| high | hypromellose | Aurolab |" in to_markdown(subject, sections)
    assert "<h3>Recalls worth watching</h3>" in to_html(subject, sections)
    assert "$" not in text.split("KEEPING IT RUNNING")[0]   # no money in the reader's part


def test_compose_handles_a_quiet_week():
    facts = {"today": TODAY, "recalls": 0, "events": 0, "flagged": [], "high": 0, "watch": 0,
             "in_shortage": [], "week": None, "recent_signals": 0, "listings_added": 0,
             "listings_left": 0, "runs": 7, "runs_ok": 7, "warnings": 0,
             "tag_review": [], "tag_review_recalls": 0, "extract_review": []}
    subject, sections = compose(facts)
    text = to_text(subject, sections)
    assert subject == "ShortageSignal weekly, Oct 19: no new recalls"
    assert "FDA reported no new recalls in the 7 days to October 19." in text
    assert "No recall this week was flagged high or watch." in text
    assert "7 of 7 pipeline runs worked this week, and every data check passed." in text


def test_workbook_has_a_sheet_per_list(tmp_path):
    path = write_workbook(gather(TODAY, _fixture(tmp_path), NO_SETTINGS), tmp_path / "out" / "needs_a_look.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Flagged this week", "Already in shortage", "Unsure tags",
                             "Drug names not used"]
    flagged = wb["Flagged this week"]
    assert flagged["A1"].value == "flag" and flagged["A2"].value == "high"
    assert wb["Drug names not used"].max_row == 2            # header and this week's one


def test_email_settings_and_message(tmp_path, monkeypatch):
    monkeypatch.delenv("DIGEST_SMTP_USER", raising=False)
    monkeypatch.delenv("DIGEST_SMTP_PASSWORD", raising=False)
    assert email_settings() is None
    monkeypatch.setenv("DIGEST_SMTP_USER", "me@example.com")
    monkeypatch.setenv("DIGEST_SMTP_PASSWORD", "app-password")
    monkeypatch.delenv("DIGEST_TO", raising=False)
    settings = email_settings()
    assert (settings["to"], settings["host"], settings["port"]) == ("me@example.com", "smtp.gmail.com", 587)

    attachment = write_workbook(gather(TODAY, _fixture(tmp_path), NO_SETTINGS), tmp_path / "list.xlsx")
    message = build_message("Subject line", "text", "<p>page</p>", attachment, settings)
    assert message["Subject"] == "Subject line" and message["To"] == "me@example.com"
    assert [part.get_filename() for part in message.iter_attachments()] == ["list.xlsx"]


def test_workbook_drops_characters_excel_cannot_store(tmp_path):
    facts = {"flagged": [{"flag": "high", "product": "Methylprednisolone\x0b 40 mg/mL vial\x02"}],
             "in_shortage": [], "tag_review": [], "extract_review": []}
    wb = load_workbook(write_workbook(facts, tmp_path / "list.xlsx"))
    assert wb["Flagged this week"]["J2"].value == "Methylprednisolone 40 mg/mL vial"


def _weekly_runs():
    """A $1.50 backfill in the first week, then $0.07 a week: one Tuesday run each week."""
    runs = [{"run_at_utc": "2026-09-15T12:00:00", "status": "success",
             "claude_cost_usd": "1.00", "extract_cost_usd": "0.50"}]
    for day in ("2026-09-22", "2026-09-29", "2026-10-06", "2026-10-13"):
        runs.append({"run_at_utc": f"{day}T12:00:00", "status": "success",
                     "claude_cost_usd": "0.03", "extract_cost_usd": "0.04"})
    runs.append({"run_at_utc": "2026-08-01T12:00:00", "status": "success",   # before buying
                 "claude_cost_usd": "9.00", "extract_cost_usd": "0"})
    return runs


def test_upkeep_estimates_credits_from_the_usual_week():
    settings = {"bought": 5.0, "since": date(2026, 9, 14), "key_expires": date(2027, 4, 13)}
    f = upkeep(_weekly_runs(), TODAY, settings)
    c = f["credits"]
    assert (c["spent"], c["left"], c["per_week"]) == (1.78, 3.22, 0.07)   # the backfill week is outvoted
    assert round(c["days_left"]) == 322
    assert f["key_days"] == 176
    lines = keeping_it_running(f)
    assert "about $3.22 left of the $5.00 bought on Sep 14, 2026" in lines[1]
    assert "about $0.07 a week, that lasts about 11 months" in lines[1]
    assert lines[2] == "The Claude API key expires on Apr 13, 2027, in 176 days."


def test_upkeep_waits_for_enough_history_and_warns():
    settings = {"bought": 5.0, "since": date(2026, 10, 13), "key_expires": date(2026, 11, 1)}
    runs = [{"run_at_utc": "2026-10-14T12:00:00", "status": "success", "claude_cost_usd": "4.40",
             "extract_cost_usd": "0", "tag_backlog": "4", "extract_waiting": "0"}]
    lines = keeping_it_running(upkeep(runs, TODAY, settings))
    assert lines[0].startswith("4 recalls are waiting for Claude")
    assert "There isn't enough history yet" in lines[1] and "Top up soon." in lines[1]
    assert "in 13 days. Make a new one soon" in lines[2]
    assert keeping_it_running({"waiting": 0})[0] == "No recalls are waiting for Claude."
