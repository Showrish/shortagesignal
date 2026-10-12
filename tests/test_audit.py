"""The hand-audit sheet and its scoring. Uses openpyxl only, no DuckDB."""
import csv

import pytest
from openpyxl import load_workbook

from shortagesignal import audit
from shortagesignal.audit import (
    PIPELINE_COLUMNS,
    YOUR_COLUMNS,
    carry_over,
    leftover_numbers,
    make_sheet,
    prefill_ndc_ok,
    read_sheet,
    score,
    write_results,
)


def pipeline_row(n, drug="", fields_missing=""):
    return {"recall_number": f"D-{n}", "report_date": "20260107", "class": "Class II",
            "firm": "Some Co", "product_description": "Some drug", "reason_for_recall": "CGMP",
            "ndcs_in_text": "", "leftover_numbers": "", "drug_found": drug, "how_found": "",
            "linked_shortage": "none",
            "fields_missing": fields_missing}


def test_sheet_has_yellow_dropdown_columns_and_a_guide(tmp_path):
    path = tmp_path / "audit_sheet.xlsx"
    make_sheet([pipeline_row(1), pipeline_row(2)], path)
    wb = load_workbook(path)
    ws = wb["Audit"]

    assert [c.value for c in ws[1]] == PIPELINE_COLUMNS + YOUR_COLUMNS
    assert ws.freeze_panes == "B2"
    first_yours = len(PIPELINE_COLUMNS) + 1
    assert ws.cell(row=1, column=first_yours).fill.start_color.rgb.endswith("FFFF00")
    assert len(ws.data_validations.dataValidation) == 4
    assert "How to fill" in wb.sheetnames


def test_scoring_turns_answers_into_caught_x_of_y(tmp_path):
    path = tmp_path / "audit_sheet.xlsx"
    make_sheet([pipeline_row(1, "heparin sodium"), pipeline_row(2, "carboplatin"),
                pipeline_row(3, fields_missing="recalling_firm"),
                pipeline_row(4, fields_missing="reason_for_recall"), pipeline_row(5)], path)
    answers = [
        ("yes", "", "yes", "none", "Sterility"),                        # right drug
        ("no", "cisplatin", "missed some", "none", "Particulates"),     # wrong drug
        ("no", "morphine sulfate", "no NDC in text", "missing field", "CGMP"),   # missed drug
        ("can't tell", "", "no NDC in text", "none", "Other"),           # a false alarm on fields
    ]                                                                    # row 5 left blank
    wb = load_workbook(path)
    ws = wb["Audit"]
    start = len(PIPELINE_COLUMNS) + 1
    for row, values in enumerate(answers, start=2):
        for offset, value in enumerate(values):
            ws.cell(row=row, column=start + offset, value=value)
    wb.save(path)

    rows = read_sheet(path)
    result = score(rows)
    assert result["audited"] == 4
    assert result["drug_precision"] == "1 of 2 (50%)"
    assert result["drug_coverage"] == "1 of 3 (33%)"
    assert result["ndc_capture"] == "1 of 2 (50%)"
    assert result["missing_flagged"] == "1 of 1 (100%)"
    assert result["false_alarms"] == 1
    assert result["labeled"] == 4

    write_results(rows, result, tmp_path, "2026-10-12")
    with open(tmp_path / "answer_key.csv", newline="") as f:
        key = list(csv.DictReader(f))
    assert [k["root_cause"] for k in key] == ["Sterility", "Particulates", "CGMP", "Other"]
    results = (tmp_path / "audit_results.md").read_text()
    assert "1 of 2 (50%)" in results
    assert "reported between 2026-01-07 and 2026-01-07" in results


def test_sample_refuses_to_overwrite_your_answers(tmp_path, monkeypatch):
    sheet = tmp_path / "audit_sheet.xlsx"
    sheet.write_text("my answers")
    monkeypatch.setattr(audit, "SHEET_PATH", sheet)
    with pytest.raises(SystemExit):
        audit.main(["sample"])
    assert sheet.read_text() == "my answers"


def test_leftover_numbers_shows_only_what_the_extractor_skipped():
    text = ("Heparin, NDC 0409-2720-01, 0409-2720-02, Lot 1234-5678, "
            "NDC 63323 540 31, exp 12/2027")
    # The two dashed NDCs were picked up. The lot number and the spaced NDC were not.
    assert leftover_numbers(text) == "1234-5678, 63323 540 31"
    assert leftover_numbers("Tablets, 30 count bottle") == ""


def test_ndc_ok_is_prefilled_only_when_nothing_is_left_over():
    assert prefill_ndc_ok({"ndcs_in_text": "00409-2720-01", "leftover_numbers": ""}) == "yes"
    assert prefill_ndc_ok({"ndcs_in_text": "", "leftover_numbers": ""}) == "no NDC in text"
    assert prefill_ndc_ok({"ndcs_in_text": "00409-2720-01", "leftover_numbers": "1234-5678"}) == ""


def test_rebuilding_the_sheet_keeps_answers_already_filled_in():
    new = [{"recall_number": "D-1", "ndc_ok": "yes"}, {"recall_number": "D-2", "ndc_ok": ""}]
    old = [{"recall_number": "D-1", "drug_ok": "no", "correct_drug": "cisplatin"},
           {"recall_number": "D-2", "ndc_ok": "missed some", "root_cause": "CGMP"}]
    rows = carry_over(new, old)
    assert rows[0] == {"recall_number": "D-1", "ndc_ok": "yes", "drug_ok": "no",
                       "correct_drug": "cisplatin"}
    assert rows[1]["ndc_ok"] == "missed some" and rows[1]["root_cause"] == "CGMP"
